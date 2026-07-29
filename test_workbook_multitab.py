"""No-spend contracts for the multi-tab workbook engine.

Every external provider and Google API surface is replaced with local fakes.
Run with ``python test_workbook_multitab.py``.
"""
from __future__ import annotations

import io
import json
import os
import tempfile
from pathlib import Path

from flask import Flask
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

import api_analyze
import intake_resolver
import job_queue
import review_contract
import review_render
import review_store
import sheets_writer
import storage_helpers
import tasks_local
import workbook_runs


class IsolatedState:
    def __init__(self, root):
        self.root = Path(root)

    def __enter__(self):
        self.old_storage = storage_helpers.STORAGE_BASE
        self.old_uploads = storage_helpers.UPLOADS_DIR
        self.old_results = storage_helpers.RESULTS_DIR
        self.old_db = job_queue.DB_PATH
        storage_helpers.STORAGE_BASE = self.root / "storage"
        storage_helpers.UPLOADS_DIR = storage_helpers.STORAGE_BASE / "uploads"
        storage_helpers.RESULTS_DIR = storage_helpers.STORAGE_BASE / "results"
        job_queue.DB_PATH = self.root / "jobs.db"
        storage_helpers.init_storage()
        job_queue.init_db()
        return self

    def __exit__(self, *_args):
        storage_helpers.STORAGE_BASE = self.old_storage
        storage_helpers.UPLOADS_DIR = self.old_uploads
        storage_helpers.RESULTS_DIR = self.old_results
        job_queue.DB_PATH = self.old_db


def make_regional_workbook(path, hidden=False, ambiguous=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("Washington", "Virginia", "New York City"):
        sheet = workbook.create_sheet(name)
        sheet.append(["Street" if ambiguous else "Property Address", "Status"])
        sheet.append([f"1 {name} St", "New"])
        sheet.append([f"2 {name} St", "Done"])
        validation = DataValidation(
            type="list", formula1='"New,Done"', allow_blank=True,
            showErrorMessage=True, errorStyle="stop",
        )
        sheet.add_data_validation(validation)
        validation.add("B2:B3")
    if hidden:
        workbook["Virginia"].sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


def test_inventory_and_grok_guardrails(root):
    path = Path(root) / "regions.xlsx"
    make_regional_workbook(path)
    snapshot = workbook_runs.inspect_xlsx(path)
    original = intake_resolver._grok_workbook_suggestions
    try:
        intake_resolver._grok_workbook_suggestions = lambda _tabs: (
            _ for _ in ()
        ).throw(AssertionError("Grok must not run for standard headers"))
        resolution = intake_resolver.resolve_workbook_schema(
            snapshot["tabs"], tasks_local.ADDRESS_VARIANTS,
        )
        assert resolution["status"] == "ready"
        assert [tab["tab"] for tab in resolution["tabs"] if tab["status"] == "selected"] == [
            "Washington", "Virginia", "New York City",
        ]
        assert [tab["row_count"] for tab in resolution["tabs"]] == [2, 2, 2]
        assert snapshot["dropdowns"][0]["inline_options"] == ["New", "Done"]

        mixed_path = Path(root) / "mixed.xlsx"
        make_regional_workbook(mixed_path)
        mixed_book = load_workbook(mixed_path)
        instructions = mixed_book.create_sheet("Instructions", 0)
        instructions.append(["Instructions"])
        instructions.append(["Upload all regional tabs without changing formatting."])
        mixed_book.save(mixed_path)
        mixed_book.close()
        mixed = intake_resolver.resolve_workbook_schema(
            workbook_runs.inspect_xlsx(mixed_path)["tabs"],
            tasks_local.ADDRESS_VARIANTS,
        )
        assert mixed["status"] == "ready"
        assert mixed["tabs"][0]["status"] == "ignored"
        assert [
            tab["tab"] for tab in mixed["tabs"] if tab["status"] == "selected"
        ] == ["Washington", "Virginia", "New York City"]

        hidden_path = Path(root) / "hidden.xlsx"
        make_regional_workbook(hidden_path, hidden=True)
        hidden = intake_resolver.resolve_workbook_schema(
            workbook_runs.inspect_xlsx(hidden_path)["tabs"],
            tasks_local.ADDRESS_VARIANTS,
        )
        assert hidden["status"] == "confirmation_required"
        assert next(tab for tab in hidden["tabs"] if tab["tab"] == "Virginia")[
            "status"
        ] == "confirmation_required"

        ambiguous_path = Path(root) / "ambiguous.xlsx"
        make_regional_workbook(ambiguous_path, ambiguous=True)
        ambiguous_tabs = workbook_runs.inspect_xlsx(ambiguous_path)["tabs"]
        intake_resolver._grok_workbook_suggestions = lambda tabs: (
            [{
                "tab": tab["tab"],
                "classification": "address",
                "address_column": "Street",
                "city_column": None,
                "state_column": None,
                "zip_column": None,
                "confidence": 0.99,
                "reason": "street-shaped values",
            } for tab in tabs],
            "",
        )
        suggested = intake_resolver.resolve_workbook_schema(
            ambiguous_tabs, tasks_local.ADDRESS_VARIANTS,
        )
        assert suggested["status"] == "confirmation_required"
        assert all(
            tab["status"] == "confirmation_required"
            and tab["suggestion"]["address_column"] == "Street"
            for tab in suggested["tabs"]
        ), "Valid Grok mappings must still require local human confirmation"

        intake_resolver._grok_workbook_suggestions = lambda tabs: (
            [{
                "tab": tabs[0]["tab"],
                "classification": "address",
                "address_column": "Street",
                "city_column": None,
                "state_column": None,
                "zip_column": None,
                "confidence": 0.99,
                "reason": "street-shaped values",
            }],
            "",
        )
        unresolved = intake_resolver.resolve_workbook_schema(
            ambiguous_tabs, tasks_local.ADDRESS_VARIANTS,
        )
        assert unresolved["status"] == "confirmation_required"
        assert all(
            tab["status"] == "unresolved" for tab in unresolved["tabs"]
        ), "Grok may not omit workbook tabs"

        intake_resolver._grok_workbook_suggestions = lambda _tabs: (
            [{
                "tab": "Invented Region",
                "classification": "address",
                "address_column": "Invented Address",
                "city_column": None,
                "state_column": None,
                "zip_column": None,
                "confidence": 1.0,
                "reason": "invented",
            }],
            "",
        )
        invented = intake_resolver.resolve_workbook_schema(
            ambiguous_tabs, tasks_local.ADDRESS_VARIANTS,
        )
        assert all(
            tab["status"] == "unresolved" for tab in invented["tabs"]
        ), "Grok may not invent tabs or columns"
    finally:
        intake_resolver._grok_workbook_suggestions = original


class FakeRequest:
    def __init__(self, payload):
        self.payload = payload

    def execute(self):
        return self.payload


class FakeConvertedSheets:
    def __init__(self, mismatch=False):
        self.mismatch = mismatch
        self.last_get = None

    def spreadsheets(self):
        return self

    def values(self):
        return self

    def get(self, **kwargs):
        if "includeGridData" in kwargs:
            options = ["New", "Changed" if self.mismatch else "Done"]
            rule = {
                "condition": {
                    "type": "ONE_OF_LIST",
                    "values": [{"userEnteredValue": value} for value in options],
                },
                "strict": True,
            }
            return FakeRequest({
                "sheets": [{
                    "data": [{
                        "rowData": [
                            {"values": [{"dataValidation": rule}]},
                            {"values": [{"dataValidation": rule}]},
                        ]
                    }]
                }]
            })
        if kwargs.get("fields", "").startswith("sheets.properties"):
            return FakeRequest({
                "sheets": [{
                    "properties": {
                        "sheetId": 10, "title": "Washington", "hidden": False,
                    }
                }]
            })
        return FakeRequest({"values": [["Property Address", "Status"]]})


def test_dropdown_conversion_fails_closed():
    snapshot = {
        "sheet_states": [{"tab": "Washington", "hidden": False}],
        "tabs": [{
            "tab": "Washington", "hidden": False,
            "headers": ["Property Address", "Status"], "row_count": 2,
        }],
        "dropdowns": [{
            "tab": "Washington",
            "ranges": ["B2:B3"],
            "formula1": '"New,Done"',
            "inline_options": ["New", "Done"],
            "resolved_options": None,
            "show_error": True,
            "error_style": "stop",
            "source_max_row": 3,
            "source_max_col": 2,
        }],
    }
    original = sheets_writer._get_services
    try:
        sheets_writer._get_services = lambda: (FakeConvertedSheets(), None)
        assert workbook_runs.verify_converted_sheet("sheet", snapshot)[
            "dropdown_cells"
        ] == 2
        sheets_writer._get_services = lambda: (
            FakeConvertedSheets(mismatch=True), None,
        )
        try:
            workbook_runs.verify_converted_sheet("sheet", snapshot)
            raise AssertionError("changed dropdown options were accepted")
        except ValueError as exc:
            assert "dropdown options" in str(exc)
    finally:
        sheets_writer._get_services = original


class FakeRangeConvertedSheets:
    def __init__(self, changed_source=False):
        self.changed_source = changed_source

    def spreadsheets(self):
        return self

    def values(self):
        return self

    def get(self, **kwargs):
        if kwargs.get("includeGridData"):
            rule = {
                "condition": {
                    "type": "ONE_OF_RANGE",
                    "values": [{"userEnteredValue": "'Options'!A1:A2"}],
                },
                "strict": True,
            }
            return FakeRequest({
                "sheets": [{
                    "data": [{
                        "rowData": [
                            {"values": [{"dataValidation": rule}]},
                            {"values": [{"dataValidation": rule}]},
                        ]
                    }]
                }]
            })
        if kwargs.get("fields", "").startswith("sheets.properties"):
            return FakeRequest({
                "sheets": [
                    {"properties": {
                        "sheetId": 10, "title": "Washington", "hidden": False,
                    }},
                    {"properties": {
                        "sheetId": 20, "title": "Options", "hidden": True,
                    }},
                ]
            })
        range_name = kwargs.get("range", "")
        if range_name == "'Washington'!1:1":
            return FakeRequest({"values": [["Property Address", "Status"]]})
        if range_name == "'Options'!1:1":
            return FakeRequest({"values": [["New"]]})
        if range_name == "'Options'!A1:A2":
            return FakeRequest({
                "values": [["New"], ["Changed" if self.changed_source else "Done"]]
            })
        return FakeRequest({"values": []})


def test_range_dropdown_values_and_unsupported_rules(root):
    snapshot = {
        "sheet_states": [
            {"tab": "Washington", "hidden": False},
            {"tab": "Options", "hidden": True},
        ],
        "tabs": [
            {
                "tab": "Washington", "hidden": False,
                "headers": ["Property Address", "Status"], "row_count": 2,
            },
            {
                "tab": "Options", "hidden": True,
                "headers": ["New"], "row_count": 1,
            },
        ],
        "dropdowns": [{
            "tab": "Washington",
            "ranges": ["B2:B3"],
            "formula1": "'Options'!$A$1:$A$2",
            "inline_options": None,
            "resolved_options": ["New", "Done"],
            "source_tab": "Options",
            "source_range": "A1:A2",
            "resolved_formula": "'Options'!A1:A2",
            "show_error": True,
            "error_style": "stop",
            "source_max_row": 3,
            "source_max_col": 2,
        }],
    }
    original = sheets_writer._get_services
    try:
        sheets_writer._get_services = lambda: (FakeRangeConvertedSheets(), None)
        assert workbook_runs.verify_converted_sheet("sheet", snapshot)[
            "dropdown_cells"
        ] == 2
        sheets_writer._get_services = lambda: (
            FakeRangeConvertedSheets(changed_source=True), None,
        )
        try:
            workbook_runs.verify_converted_sheet("sheet", snapshot)
            raise AssertionError("changed range-backed dropdown values were accepted")
        except ValueError as exc:
            assert "source values" in str(exc)
    finally:
        sheets_writer._get_services = original

    path = Path(root) / "unsupported-dropdown.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Washington"
    sheet.append(["Property Address", "Status"])
    sheet.append(["1 Main St", "New"])
    validation = DataValidation(
        type="list", formula1='INDIRECT("A1:A2")',
        showErrorMessage=True, errorStyle="stop",
    )
    sheet.add_data_validation(validation)
    validation.add("B2")
    workbook.save(path)
    workbook.close()
    try:
        workbook_runs.inspect_xlsx(path)
        raise AssertionError("unsupported dropdown source was accepted")
    except ValueError as exc:
        assert "cannot safely verify" in str(exc)


def test_conversion_mismatch_trashes_before_queue(root):
    with IsolatedState(root):
        calls = []
        originals = {
            "convert": workbook_runs._convert_local_xlsx,
            "verify": workbook_runs.verify_converted_sheet,
            "trash": workbook_runs._trash_sheet,
            "build": workbook_runs._build_queue,
        }
        try:
            workbook_runs._convert_local_xlsx = (
                lambda *_args, **_kwargs: calls.append("convert") or "converted-sheet"
            )
            workbook_runs.verify_converted_sheet = (
                lambda *_args, **_kwargs: (
                    calls.append("verify"),
                    (_ for _ in ()).throw(ValueError("dropdown mismatch")),
                )[1]
            )
            workbook_runs._trash_sheet = (
                lambda sheet_id: calls.append(f"trash:{sheet_id}") or True
            )
            workbook_runs._build_queue = lambda _run: (
                calls.append("queue"),
                (_ for _ in ()).throw(AssertionError("queue must not start")),
            )[1]
            run = {
                "schema_version": 2,
                "run_id": "w-mismatch",
                "source_name": "regions.xlsx",
                "source_blob": "workbook_runs/w-mismatch/source.xlsx",
                "xlsx_snapshot": {},
                "created_at": "now",
            }
            stopped = workbook_runs._convert_and_prepare(run)
            assert stopped["status"] == "needs_attention"
            assert stopped["preflight_failed"] is True
            assert stopped["converted_copy_trashed"] is True
            assert calls == ["convert", "verify", "trash:converted-sheet"]
            assert job_queue.get_job_status("w-mismatch") is None
        finally:
            workbook_runs._convert_local_xlsx = originals["convert"]
            workbook_runs.verify_converted_sheet = originals["verify"]
            workbook_runs._trash_sheet = originals["trash"]
            workbook_runs._build_queue = originals["build"]


def test_queue_chunking_duplicate_reuse_and_atomic_approval(root):
    with IsolatedState(root):
        old_threshold = os.environ.get("WORKBOOK_AUTO_APPROVAL_ROWS")
        old_min = os.environ.get("ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS")
        old_max = os.environ.get("ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS")
        original_read = sheets_writer.read_bound_sheet_mappings
        original_columns = sheets_writer.ensure_review_columns
        try:
            os.environ["WORKBOOK_AUTO_APPROVAL_ROWS"] = "2"
            os.environ["ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS"] = "0.01"
            os.environ["ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS"] = "0.02"
            sheets_writer.ensure_review_columns = lambda binding, *_args, **_kwargs: (
                binding.update({"colmap": {"HVAC Systems": 1}}) or binding
            )
            sheets_writer.read_bound_sheet_mappings = lambda _url, _mappings: [
                {
                    "spreadsheet_id": "sheet", "grid_id": 10, "tab": "Washington",
                    "headers": ["Address"], "row_numbers": [2, 3],
                    "targets": [
                        {"source_key": "g10:r2", "grid_id": 10, "tab": "Washington",
                         "source_row": 2, "address": "1 Main St"},
                        {"source_key": "g10:r3", "grid_id": 10, "tab": "Washington",
                         "source_row": 3, "address": "2 Main St"},
                    ],
                },
                {
                    "spreadsheet_id": "sheet", "grid_id": 20, "tab": "Virginia",
                    "headers": ["Address"], "row_numbers": [2],
                    "targets": [
                        {"source_key": "g20:r2", "grid_id": 20, "tab": "Virginia",
                         "source_row": 2, "address": " 1  MAIN ST "},
                    ],
                },
            ]
            run = {
                "run_id": "w-test",
                "sheet_url": "https://docs.google.com/spreadsheets/d/sheet/edit",
                "tabs": [
                    {"tab": "Washington", "status": "selected",
                     "mapping": {"tab": "Washington", "address_column": "Address"}},
                    {"tab": "Virginia", "status": "selected",
                     "mapping": {"tab": "Virginia", "address_column": "Address"}},
                ],
                "created_at": "now",
            }
            built = workbook_runs._build_queue(run)
            assert built["row_count"] == 3
            assert built["analysis_count"] == 2
            assert built["status"] == "approval_required"
            approved = workbook_runs._enqueue(built, approved=True)
            assert approved["status"] == "queued"
            assert job_queue.get_job_status("w-test")["total"] == 2
            assert job_queue.get_monthly_usage() == 3

            job_queue.update_job_status("w-test", status="processing")
            assert job_queue.recover_interrupted_workbook_jobs() == 1
            assert job_queue.get_job_status("w-test")["status"] == "queued"
            job_queue.update_job_status("w-test", status="failed")
            workbook_runs.mutate_run(
                "w-test",
                lambda current: current.update({
                    "status": "needs_attention",
                    "error": "simulated worker stop",
                }),
            )
            retried = workbook_runs.retry("w-test")
            assert retried["status"] == "queued"
            assert job_queue.get_job_status("w-test")["status"] == "queued"
            assert job_queue.get_monthly_usage() == 3
            workbook_runs.update_review_progress(
                "w-test", 3, 3, writeback_failures=1, needs_attention=0,
            )
            assert workbook_runs.load_run("w-test")["status"] == "needs_attention"
            workbook_runs.update_review_progress(
                "w-test", 3, 3, writeback_failures=0, needs_attention=0,
            )
            assert workbook_runs.load_run("w-test")["status"] == "complete"
        finally:
            sheets_writer.read_bound_sheet_mappings = original_read
            sheets_writer.ensure_review_columns = original_columns
            for name, value in (
                ("WORKBOOK_AUTO_APPROVAL_ROWS", old_threshold),
                ("ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS", old_min),
                ("ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS", old_max),
            ):
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value


def test_cost_configuration_is_optional_but_validated():
    old_values = {
        name: os.environ.get(name)
        for name in (
            "ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS",
            "ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS",
        )
    }
    try:
        os.environ.pop("ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS", None)
        os.environ.pop("ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS", None)
        status = workbook_runs.cost_configuration_status()
        assert status["configured"] is False
        assert len(status["missing_configuration"]) == 2
        assert workbook_runs._cost_estimate(300) is None

        for low, high in (
            ("nan", "0.02"),
            ("0.01", "inf"),
            ("-0.01", "0.02"),
            ("0.03", "0.02"),
        ):
            os.environ["ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS"] = low
            os.environ["ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS"] = high
            status = workbook_runs.cost_configuration_status()
            assert status["configured"] is False
            assert len(status["invalid_configuration"]) == 2
            assert workbook_runs._cost_estimate(300) is None

        os.environ["ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS"] = "0.01"
        os.environ["ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS"] = "0.02"
        status = workbook_runs.cost_configuration_status()
        assert status["configured"] is True
        assert status["_rates"] == (0.01, 0.02)
        assert workbook_runs._cost_estimate(300) == {
            "currency": "USD",
            "minimum": 3.0,
            "maximum": 6.0,
            "basis": "configured_per_unique_address",
        }
    finally:
        for name, value in old_values.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value


def test_large_workbook_approval_without_cost_rates(root):
    with IsolatedState(root):
        old_threshold = os.environ.get("WORKBOOK_AUTO_APPROVAL_ROWS")
        old_values = {
            name: os.environ.get(name)
            for name in (
                "ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS",
                "ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS",
            )
        }
        try:
            os.environ["WORKBOOK_AUTO_APPROVAL_ROWS"] = "2"
            for name in old_values:
                os.environ.pop(name, None)
            run = {
                "run_id": "w-no-cost",
                "status": "approval_required",
                "analysis_count": 3,
                "row_count": 3,
                "cost_estimate": None,
                "created_at": "now",
            }
            workbook_runs.save_run(run)
            approved = workbook_runs.approve("w-no-cost")
            assert approved["status"] == "queued"
            assert approved["approved_at"]
            assert approved["reserved_address_count"] == 3
            assert job_queue.get_job_status("w-no-cost")["total"] == 3
        finally:
            if old_threshold is None:
                os.environ.pop("WORKBOOK_AUTO_APPROVAL_ROWS", None)
            else:
                os.environ["WORKBOOK_AUTO_APPROVAL_ROWS"] = old_threshold
            for name, value in old_values.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value


def test_costless_approval_ui_is_enabled():
    source = Path("templates/workbook_approval.html").read_text(encoding="utf-8")
    assert "Approval will remain blocked" not in source
    assert "if not run.cost_estimate %}disabled" not in source
    assert "No cost estimate is configured" in source


def test_default_hundred_address_chunks(root):
    with IsolatedState(root):
        original_read = sheets_writer.read_bound_sheet_mappings
        original_columns = sheets_writer.ensure_review_columns
        old_chunk = os.environ.get("WORKBOOK_CHUNK_ROWS")
        try:
            os.environ["WORKBOOK_CHUNK_ROWS"] = "100"
            sheets_writer.ensure_review_columns = (
                lambda binding, *_args, **_kwargs:
                binding.update({"colmap": {"HVAC Systems": 1}}) or binding
            )
            targets = [
                {
                    "source_key": f"sabc:g10:r{row}",
                    "grid_id": 10,
                    "tab": "Washington",
                    "source_row": row,
                    "address": f"{row} Main St",
                }
                for row in range(2, 207)
            ]
            sheets_writer.read_bound_sheet_mappings = lambda *_args, **_kwargs: [{
                "spreadsheet_id": "sheet",
                "grid_id": 10,
                "tab": "Washington",
                "headers": ["Address"],
                "row_numbers": list(range(2, 207)),
                "targets": targets,
            }]
            run = {
                "run_id": "w-chunks",
                "sheet_url": "https://docs.google.com/spreadsheets/d/sheet/edit",
                "tabs": [{
                    "tab": "Washington",
                    "status": "selected",
                    "mapping": {
                        "tab": "Washington", "address_column": "Address",
                    },
                }],
                "created_at": "now",
            }
            built = workbook_runs._build_queue(run)
            assert built["row_count"] == 205
            assert built["analysis_count"] == 205
            assert [len(chunk["items"]) for chunk in built["chunks"]] == [100, 100, 5]
        finally:
            sheets_writer.read_bound_sheet_mappings = original_read
            sheets_writer.ensure_review_columns = original_columns
            if old_chunk is None:
                os.environ.pop("WORKBOOK_CHUNK_ROWS", None)
            else:
                os.environ["WORKBOOK_CHUNK_ROWS"] = old_chunk


def test_resume_skips_checkpointed_rows(root):
    csv_path = Path(root) / "chunk.csv"
    csv_path.write_text("Address\n1 Main St\n2 Main St\n", encoding="utf-8")
    old_key = os.environ.get("MAPBOX_API_KEY")
    original = tasks_local._process_one_address
    calls = []
    partials = []
    try:
        os.environ["MAPBOX_API_KEY"] = "mock"

        def fake_process(row, index, *_args, **_kwargs):
            calls.append(int(index))
            address = str(row["Address"])
            return (
                {"address": address, "verdict": "not_detected", "error": ""},
                {"Address": address},
            )

        tasks_local._process_one_address = fake_process
        result = tasks_local.process_address_list(
            str(csv_path), "resume-test",
            progress_cb=lambda *_args: None,
            should_cancel=lambda: False,
            upload_file=lambda _local, blob: blob,
            make_signed_url=lambda blob: f"/files/{blob}",
            write_partial_result=lambda data: partials.append(data),
            concurrency=1,
            resume_results={
                0: {
                    "web_entry": {
                        "address": "1 Main St", "verdict": "not_detected", "error": "",
                    },
                    "csv_row": {"Address": "1 Main St"},
                }
            },
        )
        assert calls == [1]
        assert [entry["address"] for entry in result["web_results"]] == [
            "1 Main St", "2 Main St",
        ]
        assert any(
            next(
                state for state in partial["address_states"]
                if state["index"] == 1
            )["state"] == "analyzing"
            for partial in partials
        )
        assert partials[-1]["address_states"] == [
            {
                "index": 0,
                "address": "1 Main St",
                "state": "complete",
                "message": "Machine analysis finished",
                "error": "",
            },
            {
                "index": 1,
                "address": "2 Main St",
                "state": "complete",
                "message": "Machine analysis finished",
                "error": "",
            },
        ]
    finally:
        tasks_local._process_one_address = original
        if old_key is None:
            os.environ.pop("MAPBOX_API_KEY", None)
        else:
            os.environ["MAPBOX_API_KEY"] = old_key


def test_grouped_review_and_exact_tab_writeback():
    entries = [
        {"row_id": "g10:r2", "source_tab": "Washington", "source_row": 2,
         "address": "1 Main St", "verdict": "not_detected"},
        {"row_id": "g20:r2", "source_tab": "Virginia", "source_row": 2,
         "address": "2 Main St", "verdict": "not_detected"},
    ]
    html = review_render.build_review_page(entries, "w-test")
    assert "<h2>Washington</h2>" in html and "<h2>Virginia</h2>" in html
    assert 'data-rid="g10:r2"' in html and 'data-rid="g20:r2"' in html
    assert "choose both Optimizer Fit and Periscope Fit" in html
    assert "Optimizer Fit:" in html and "Periscope Fit:" in html

    class FakeSheets:
        def __init__(self):
            self.updates = []

        def spreadsheets(self):
            return self

        def values(self):
            return self

        def batchUpdate(self, **kwargs):
            self.updates.append(kwargs)
            return FakeRequest({})

    fake = FakeSheets()
    original = sheets_writer._get_services
    try:
        sheets_writer._get_services = lambda: (fake, None)
        for tab, grid in (("Washington", 10), ("Virginia", 20)):
            assert sheets_writer.write_decision_source(
                {
                    "spreadsheet_id": "sheet", "grid_id": grid, "tab": tab,
                    "colmap": {
                        "HVAC Systems": 2,
                        "Optimizer Fit": 3,
                        "Periscope Fit": 4,
                    },
                },
                2, "None", optimizer_fit="Good", periscope_fit="Bad",
            )
        ranges = [
            item["range"]
            for update in fake.updates
            for item in update["body"]["data"]
        ]
        assert any("'Washington'!C2" == value for value in ranges)
        assert any("'Virginia'!C2" == value for value in ranges)
        assert any("'Washington'!D2" == value for value in ranges)
        assert any("'Virginia'!D2" == value for value in ranges)
        assert any("'Washington'!E2" == value for value in ranges)
        assert any("'Virginia'!E2" == value for value in ranges)
        assert sheets_writer.write_source_values(
            {
                "spreadsheet_id": "sheet", "grid_id": 20, "tab": "Virginia",
                "headers": ["Property Address"],
            },
            2,
            {"Property Address": "Corrected Virginia Address"},
        )
        assert any(
            item["range"] == "'Virginia'!A2"
            for update in fake.updates
            for item in update["body"]["data"]
        )
    finally:
        sheets_writer._get_services = original


def test_dual_fit_columns_are_reused_or_appended_once():
    class FakeSheets:
        def __init__(self):
            self.value_updates = []
            self.grid_updates = []

        def spreadsheets(self):
            return self

        def values(self):
            return self

        def update(self, **kwargs):
            self.value_updates.append(kwargs)
            return FakeRequest({})

        def batchUpdate(self, **kwargs):
            self.grid_updates.append(kwargs)
            return FakeRequest({})

    original = sheets_writer._get_services
    try:
        existing = FakeSheets()
        sheets_writer._get_services = lambda: (existing, None)
        binding = {
            "spreadsheet_id": "sheet",
            "grid_id": 10,
            "tab": "Washington",
            "row_numbers": [2, 3],
        }
        sheets_writer.ensure_review_columns(
            binding,
            [
                "Property Address",
                "HVAC Systems",
                "Optimizer Fit",
                "Periscope Fit",
                "Notes",
            ],
            review_render.HVAC_SYSTEMS + [review_render.NONE_OPTION],
            review_contract.DUAL_FIT_OPTIONS,
        )
        assert binding["colmap"] == {
            "HVAC Systems": 1,
            "Optimizer Fit": 2,
            "Periscope Fit": 3,
            "Notes": 4,
        }
        assert existing.value_updates == []
        assert existing.grid_updates == []

        missing = FakeSheets()
        sheets_writer._get_services = lambda: (missing, None)
        binding = {
            "spreadsheet_id": "sheet",
            "grid_id": 20,
            "tab": "Virginia",
            "row_numbers": [2, 3],
        }
        sheets_writer.ensure_review_columns(
            binding,
            ["Property Address", "HVAC Systems", "Notes"],
            review_render.HVAC_SYSTEMS + [review_render.NONE_OPTION],
            review_contract.DUAL_FIT_OPTIONS,
        )
        assert missing.value_updates[0]["body"]["values"] == [[
            "Optimizer Fit", "Periscope Fit",
        ]]
        validation_requests = [
            request["setDataValidation"]
            for update in missing.grid_updates
            for request in update["body"]["requests"]
            if "setDataValidation" in request
        ]
        assert len(validation_requests) == 2
        for validation in validation_requests:
            options = [
                item["userEnteredValue"]
                for item in validation["rule"]["condition"]["values"]
            ]
            assert options == ["Good", "Bad", "Not Sure"]
            assert validation["rule"]["strict"] is True
    finally:
        sheets_writer._get_services = original


def test_converted_copy_review_answers_are_cleared_without_rule_changes():
    class FakeSheets:
        def __init__(self):
            self.clears = []

        def spreadsheets(self):
            return self

        def values(self):
            return self

        def batchClear(self, **kwargs):
            self.clears.append(kwargs)
            return FakeRequest({})

    original = sheets_writer._get_services
    try:
        fake = FakeSheets()
        sheets_writer._get_services = lambda: (fake, None)
        binding = {
            "spreadsheet_id": "sheet",
            "grid_id": 20,
            "tab": "Road Trip Chaos",
            "row_numbers": [2, 3, 11],
            "colmap": {
                "HVAC Systems": 8,
                "Optimizer Fit": 9,
                "Periscope Fit": 10,
                "Notes": 11,
            },
        }
        assert sheets_writer.clear_review_answers(binding)
        assert fake.clears == [{
            "spreadsheetId": "sheet",
            "body": {
                "ranges": [
                    "'Road Trip Chaos'!I2:I11",
                    "'Road Trip Chaos'!J2:J11",
                    "'Road Trip Chaos'!K2:K11",
                ],
            },
        }]
    finally:
        sheets_writer._get_services = original


def test_corrected_rerun_uses_exact_multitab_source(root):
    with IsolatedState(root):
        batch_id = "w-rerun-source"
        entry = {
            "row_id": "sabc:g20:r2",
            "i": "sabc:g20:r2",
            "source_grid_id": 20,
            "source_tab": "Virginia",
            "source_row": 2,
            "address": "Bad Address",
            "error": "geocode failed",
        }
        review_store.save_batch(
            batch_id,
            "Regions",
            [entry],
            sheet_url="https://docs.google.com/spreadsheets/d/sheet/edit",
            table_headers=["Address", "Row_ID"],
            table_rows=[["Bad Address", "sabc:g20:r2"]],
            sheet_bindings=[{
                "spreadsheet_id": "sheet",
                "grid_id": 20,
                "tab": "Virginia",
                "headers": ["Property Address"],
                "intake_mapping": {"address_column": "Property Address"},
            }],
            schema_version=2,
        )
        old_key = os.environ.get("ANALYZE_API_KEY")
        original_analyze = api_analyze._analyze_one
        original_enabled = sheets_writer.enabled
        original_source_write = sheets_writer.write_source_values
        original_legacy_write = sheets_writer.write_row_values
        writes = []
        try:
            os.environ["ANALYZE_API_KEY"] = "test"
            api_analyze._analyze_one = lambda address, **_kwargs: {
                "address": address,
                "verdict": "not_detected",
                "error": "",
            }
            sheets_writer.enabled = lambda: True
            sheets_writer.write_source_values = (
                lambda binding, row, values:
                writes.append((binding["tab"], row, values)) or True
            )
            sheets_writer.write_row_values = lambda *_args, **_kwargs: (
                _ for _ in ()
            ).throw(AssertionError("legacy first-tab writer must not run"))
            app = Flask("workbook-rerun-source-test")
            app.register_blueprint(api_analyze.api)
            response = app.test_client().post(
                f"/api/batch/{batch_id}/rerun",
                json={"rows": [{
                    "row_id": "sabc:g20:r2",
                    "address": "2 Corrected Virginia Ave",
                }]},
                headers={"X-API-Key": "test"},
            )
            assert response.status_code == 200
            assert writes == [(
                "Virginia",
                2,
                {"Property Address": "2 Corrected Virginia Ave"},
            )]
        finally:
            api_analyze._analyze_one = original_analyze
            sheets_writer.enabled = original_enabled
            sheets_writer.write_source_values = original_source_write
            sheets_writer.write_row_values = original_legacy_write
            if old_key is None:
                os.environ.pop("ANALYZE_API_KEY", None)
            else:
                os.environ["ANALYZE_API_KEY"] = old_key


def test_legacy_run_file_fails_closed_on_multiple_tabs(root):
    path = Path(root) / "multi.xlsx"
    make_regional_workbook(path)
    old_key = os.environ.get("ANALYZE_API_KEY")
    try:
        os.environ["ANALYZE_API_KEY"] = "test"
        app = Flask("workbook-legacy-test")
        app.register_blueprint(api_analyze.api)
        with path.open("rb") as handle:
            response = app.test_client().post(
                "/api/run-file",
                data={"file": (io.BytesIO(handle.read()), "multi.xlsx")},
                headers={"X-API-Key": "test"},
            )
        assert response.status_code == 409
        assert response.json["status"] == "workbook_endpoint_required"
    finally:
        if old_key is None:
            os.environ.pop("ANALYZE_API_KEY", None)
        else:
            os.environ["ANALYZE_API_KEY"] = old_key


def test_versioned_async_api_contract():
    old_values = {
        name: os.environ.get(name) for name in (
            "ANALYZE_API_KEY", "MULTI_TAB_WORKBOOK_ENABLED",
            "MULTI_TAB_API_ENABLED",
        )
    }
    original_prepare = workbook_runs.prepare_local_file
    try:
        os.environ["ANALYZE_API_KEY"] = "test"
        os.environ["MULTI_TAB_WORKBOOK_ENABLED"] = "true"
        os.environ["MULTI_TAB_API_ENABLED"] = "true"
        workbook_runs.prepare_local_file = lambda *_args, **_kwargs: {
            "schema_version": 2,
            "run_id": "w-contract",
            "status": "queued",
            "source_name": "regions.csv",
            "created_at": "now",
            "updated_at": "now",
            "tabs": [{
                "tab": "Uploaded file", "headers": ["Address"],
                "status": "selected", "row_count": 2,
                "mapping": {"address_column": "Address"},
            }],
            "row_count": 2,
            "analysis_count": 2,
            "sheet_url": "https://docs.google.com/spreadsheets/d/sheet/edit",
        }
        app = Flask("workbook-v2-test")
        app.register_blueprint(api_analyze.api)
        response = app.test_client().post(
            "/api/v2/workbook-runs",
            data={"file": (io.BytesIO(b"Address\n1 Main St\n"), "regions.csv")},
            headers={"X-API-Key": "test"},
        )
        assert response.status_code == 202
        assert response.json["schema_version"] == 2
        assert response.json["run_id"] == "w-contract"
        assert response.json["status_url"].endswith(
            "/api/v2/workbook-runs/w-contract"
        )
        assert response.json["retry_url"].endswith(
            "/api/v2/workbook-runs/w-contract/retry"
        )
        assert response.json["tab_inventory"][0]["headers"] == ["Address"]
    finally:
        workbook_runs.prepare_local_file = original_prepare
        for name, value in old_values.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value


def test_n8n_async_workflow_contract():
    workflow = json.loads(
        Path("n8n/parity_workbook_async.workflow.json").read_text(
            encoding="utf-8",
        )
    )
    nodes = {node["name"]: node for node in workflow["nodes"]}
    create = nodes["Create Workbook Run"]
    poll = nodes["Poll Workbook Status"]
    assert create["parameters"]["url"].endswith("/api/v2/workbook-runs")
    assert create["parameters"]["authentication"] == "genericCredentialType"
    assert poll["parameters"]["authentication"] == "genericCredentialType"
    assert "headerParameters" not in create["parameters"]
    assert "headerParameters" not in poll["parameters"]
    assert create["credentials"]["httpHeaderAuth"]["name"] == (
        "Parity Render API Key"
    )
    assert workflow["connections"]["Poll Workbook Status"]["main"][0][0][
        "node"
    ] == "Still Running?"
    serialized = json.dumps(workflow)
    assert "REPLACE_WITH_ANALYZE_API_KEY" not in serialized
    assert "Grok normalize" not in serialized


if __name__ == "__main__":
    with tempfile.TemporaryDirectory() as temp_dir:
        test_inventory_and_grok_guardrails(temp_dir)
        test_range_dropdown_values_and_unsupported_rules(temp_dir)
        test_conversion_mismatch_trashes_before_queue(temp_dir)
        test_queue_chunking_duplicate_reuse_and_atomic_approval(temp_dir)
        test_cost_configuration_is_optional_but_validated()
        test_large_workbook_approval_without_cost_rates(temp_dir)
        test_default_hundred_address_chunks(temp_dir)
        test_resume_skips_checkpointed_rows(temp_dir)
        test_corrected_rerun_uses_exact_multitab_source(temp_dir)
        test_legacy_run_file_fails_closed_on_multiple_tabs(temp_dir)
    test_dropdown_conversion_fails_closed()
    test_grouped_review_and_exact_tab_writeback()
    test_dual_fit_columns_are_reused_or_appended_once()
    test_converted_copy_review_answers_are_cleared_without_rule_changes()
    test_versioned_async_api_contract()
    test_costless_approval_ui_is_enabled()
    test_n8n_async_workflow_contract()
    print("OK: multi-tab workbook contracts hold without paid or network calls.")
