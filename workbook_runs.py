"""Durable multi-tab workbook intake and run orchestration.

One workbook becomes one Google Sheet, one logical analysis run, and one review
batch.  Every analyzed row keeps an opaque grid-id/physical-row source key so
duplicate row numbers on different tabs can never collide during write-back.
"""
from __future__ import annotations

import collections
import csv
import hashlib
import io
import json
import logging
import math
import os
import re
import tempfile
import threading
import uuid
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

import intake_resolver
import sheets_writer
from job_queue import (
    enqueue_job_with_usage_reservation,
    get_job_status,
    requeue_failed_workbook_job,
)
from storage_helpers import (
    delete_file,
    get_file_path,
    read_json,
    upload_file,
    write_json,
)


log = logging.getLogger("workbook_runs")
SCHEMA_VERSION = 2
RUN_ROOT = "workbook_runs"
_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_locks_guard = threading.Lock()
_locks: dict[str, threading.RLock] = {}
_metrics_lock = threading.Lock()
_metrics = collections.Counter()
_COST_RATE_ENV_NAMES = (
    "ANALYZE_ESTIMATED_MIN_COST_PER_ADDRESS",
    "ANALYZE_ESTIMATED_MAX_COST_PER_ADDRESS",
)


def enabled() -> bool:
    raw = os.getenv("MULTI_TAB_WORKBOOK_ENABLED", "false")
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def surface_enabled(surface: str, default: bool = False) -> bool:
    if not enabled():
        return False
    raw = os.getenv(f"MULTI_TAB_{surface.upper()}_ENABLED")
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def record_metric(name: str, count: int = 1) -> None:
    with _metrics_lock:
        _metrics[str(name)] += int(count)


def metrics_snapshot() -> dict[str, int]:
    with _metrics_lock:
        return dict(_metrics)


def auto_approval_threshold() -> int:
    try:
        return max(1, int(os.getenv("WORKBOOK_AUTO_APPROVAL_ROWS", "250")))
    except ValueError:
        return 250


def chunk_size() -> int:
    try:
        return max(1, min(500, int(os.getenv("WORKBOOK_CHUNK_ROWS", "100"))))
    except ValueError:
        return 100


def _lock_for(run_id: str) -> threading.RLock:
    with _locks_guard:
        lock = _locks.get(str(run_id))
        if lock is None:
            lock = threading.RLock()
            _locks[str(run_id)] = lock
        return lock


def _run_path(run_id: str) -> str:
    return f"{RUN_ROOT}/{run_id}/manifest.json"


def _utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def save_run(run: dict[str, Any]) -> dict[str, Any]:
    run["updated_at"] = _utcnow()
    with _lock_for(run["run_id"]):
        write_json(_run_path(run["run_id"]), run)
    return run


def load_run(run_id: str) -> dict[str, Any] | None:
    return read_json(_run_path(str(run_id)))


def mutate_run(run_id: str, updater) -> dict[str, Any] | None:
    with _lock_for(run_id):
        run = load_run(run_id)
        if not run:
            return None
        updater(run)
        run["updated_at"] = _utcnow()
        write_json(_run_path(run_id), run)
        return run


def list_runs() -> list[dict[str, Any]]:
    base = get_file_path(RUN_ROOT)
    if not base.exists():
        return []
    runs = []
    for path in base.glob("*/manifest.json"):
        try:
            run = read_json(f"{RUN_ROOT}/{path.parent.name}/manifest.json")
        except (OSError, TypeError, ValueError, json.JSONDecodeError):
            log.exception("Skipping unreadable workbook manifest %s", path)
            continue
        if run:
            runs.append(public_summary(run))
    return sorted(runs, key=lambda item: item.get("created_at", ""), reverse=True)


def cost_configuration_status() -> dict[str, Any]:
    """Return a non-secret readiness receipt for large-workbook approval."""
    missing = [
        name for name in _COST_RATE_ENV_NAMES
        if not os.environ.get(name, "").strip()
    ]
    if missing:
        return {
            "configured": False,
            "missing_configuration": missing,
            "invalid_configuration": [],
        }
    try:
        low = float(os.environ[_COST_RATE_ENV_NAMES[0]])
        high = float(os.environ[_COST_RATE_ENV_NAMES[1]])
    except (TypeError, ValueError):
        return {
            "configured": False,
            "missing_configuration": [],
            "invalid_configuration": list(_COST_RATE_ENV_NAMES),
        }
    if not math.isfinite(low) or not math.isfinite(high) or low < 0 or high < low:
        return {
            "configured": False,
            "missing_configuration": [],
            "invalid_configuration": list(_COST_RATE_ENV_NAMES),
        }
    return {
        "configured": True,
        "missing_configuration": [],
        "invalid_configuration": [],
        "_rates": (low, high),
    }


def _cost_estimate(address_count: int) -> dict[str, Any] | None:
    configuration = cost_configuration_status()
    if not configuration["configured"]:
        return None
    low, high = configuration["_rates"]
    return {
        "currency": "USD",
        "minimum": round(address_count * low, 2),
        "maximum": round(address_count * high, 2),
        "basis": "configured_per_unique_address",
    }


def public_summary(run: dict[str, Any]) -> dict[str, Any]:
    """Return run metadata without source cells, addresses, or private blob paths."""
    tabs = []
    for tab in run.get("tabs", []):
        mapping = tab.get("mapping") or tab.get("suggestion") or {}
        tabs.append({
            "tab": tab.get("tab", ""),
            "hidden": bool(tab.get("hidden")),
            "headers": [str(value) for value in tab.get("headers", [])],
            "status": tab.get("status", ""),
            "reason": tab.get("reason", ""),
            "row_count": int(tab.get("row_count") or 0),
            "address_column": mapping.get("address_column"),
            "city_column": mapping.get("city_column"),
            "state_column": mapping.get("state_column"),
            "zip_column": mapping.get("zip_column"),
            "suggested_classification": (
                (tab.get("suggestion") or {}).get("classification")
            ),
        })
    return {
        "schema_version": run.get("schema_version", SCHEMA_VERSION),
        "run_id": run.get("run_id"),
        "status": run.get("status"),
        "source_name": run.get("source_name", ""),
        "created_at": run.get("created_at"),
        "updated_at": run.get("updated_at"),
        "tabs": tabs,
        "tab_count": sum(tab["status"] == "selected" for tab in tabs),
        "row_count": int(run.get("row_count") or 0),
        "analysis_count": int(run.get("analysis_count") or 0),
        "reserved_address_count": int(run.get("reserved_address_count") or 0),
        "completed_rows": int(run.get("completed_rows") or 0),
        "reviewed_rows": int(run.get("reviewed_rows") or 0),
        "needs_attention": int(run.get("needs_attention") or 0),
        "writeback_failures": int(run.get("writeback_failures") or 0),
        "approval_required": run.get("status") == "approval_required",
        "confirmation_required": bool(run.get("confirmation_required")),
        "preflight_failed": bool(run.get("preflight_failed")),
        "cost_estimate": run.get("cost_estimate"),
        "sheet_url": run.get("sheet_url", ""),
        "review_url": run.get("review_url", ""),
        "error": run.get("error", ""),
    }


def _safe_name(filename: str) -> str:
    name = Path(filename or "workbook.xlsx").name
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", name)[:180] or "workbook.xlsx"


def _inline_list(formula: str | None) -> list[str] | None:
    text = str(formula or "").strip()
    if len(text) >= 2 and text[0] == text[-1] == '"':
        return [part.strip() for part in text[1:-1].split(",")]
    return None


def _normalized_range_formula(value: str | None) -> str:
    text = str(value or "").strip().lstrip("=")
    return re.sub(r"[\s$']", "", text).casefold()


def _direct_range_reference(
    formula: str | None, current_tab: str,
) -> tuple[str, str] | None:
    text = str(formula or "").strip().lstrip("=")
    if not text:
        return None
    match = re.fullmatch(
        r"(?:(?:'((?:[^']|'')+)'|([^!]+))!)?"
        r"(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)",
        text,
    )
    if not match:
        return None
    tab_name = (match.group(1) or match.group(2) or current_tab).replace("''", "'")
    return tab_name, match.group(3).replace("$", "")


def _resolved_list_source(workbook, current_sheet, formula: str | None):
    """Resolve one static Excel list source into a tab, range, and exact values.

    Dynamic formulas and multi-area names cannot be proven equivalent after
    Excel-to-Sheets conversion, so callers fail closed on those rules.
    """
    text = str(formula or "").strip().lstrip("=")
    if not text or _inline_list(formula) is not None:
        return None
    reference = _direct_range_reference(formula, current_sheet.title)
    if reference is None:
        defined = workbook.defined_names.get(text)
        if defined is None:
            return None
        try:
            destinations = list(defined.destinations)
        except (AttributeError, TypeError, ValueError):
            return None
        if len(destinations) != 1:
            return None
        tab_name, coordinate = destinations[0]
        tab_name = str(tab_name).replace("''", "'")
        coordinate = str(coordinate).replace("$", "")
        try:
            range_boundaries(coordinate)
        except ValueError:
            return None
        reference = tab_name, coordinate
    tab_name, coordinate = reference
    if tab_name not in workbook.sheetnames:
        return None
    min_col, min_row, max_col, max_row = range_boundaries(coordinate)
    cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
    source_limit = max(
        1, int(os.getenv("WORKBOOK_MAX_DROPDOWN_SOURCE_CELLS", "10000"))
    )
    if cell_count > source_limit:
        raise ValueError(
            f"Dropdown source range exceeds the configured {source_limit} cell limit"
        )
    cells = workbook[tab_name][coordinate]
    values = []
    for row in cells:
        for cell in row:
            if getattr(cell, "data_type", "") == "f":
                raise ValueError(
                    f"Dropdown source '{tab_name}!{coordinate}' contains formulas "
                    "whose displayed values cannot be safely fingerprinted"
                )
            if cell.value is not None:
                values.append(str(cell.value).strip())
    return {
        "tab": tab_name,
        "range": coordinate,
        "formula": f"{_quoted_tab(tab_name)}!{coordinate}",
        "options": values,
    }


def inspect_xlsx(path: str | Path) -> dict[str, Any]:
    """Read workbook structure, bounded intake samples, and dropdown contracts."""
    with open(path, "rb") as source:
        workbook_bytes = io.BytesIO(source.read())
    try:
        with zipfile.ZipFile(workbook_bytes) as archive:
            expanded_size = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as exc:
        workbook_bytes.close()
        raise ValueError("The uploaded file is not a valid .xlsx workbook") from exc
    expanded_limit = max(
        1, int(os.getenv("WORKBOOK_MAX_UNCOMPRESSED_BYTES", str(250 * 1024 * 1024)))
    )
    if expanded_size > expanded_limit:
        workbook_bytes.close()
        raise ValueError(
            f"Workbook expands beyond the configured {expanded_limit} byte limit"
        )
    workbook_bytes.seek(0)
    workbook = load_workbook(
        filename=workbook_bytes, read_only=False, data_only=False,
    )
    tabs = []
    dropdowns = []
    sheet_states = []
    max_rows_allowed = max(1, int(os.getenv("WORKBOOK_MAX_ROWS_PER_TAB", "100000")))
    max_cols_allowed = max(1, int(os.getenv("WORKBOOK_MAX_COLUMNS", "500")))
    max_cells_allowed = max(1, int(os.getenv("WORKBOOK_MAX_USED_CELLS", "5000000")))
    used_cells = 0
    try:
        for sheet in workbook.worksheets:
            max_col = max(1, sheet.max_column or 1)
            max_row = max(1, sheet.max_row or 1)
            used_cells += max_col * max_row
            if max_row > max_rows_allowed:
                raise ValueError(
                    f"Tab '{sheet.title}' exceeds the configured {max_rows_allowed} row limit"
                )
            if max_col > max_cols_allowed:
                raise ValueError(
                    f"Tab '{sheet.title}' exceeds the configured {max_cols_allowed} column limit"
                )
            if used_cells > max_cells_allowed:
                raise ValueError(
                    f"Workbook exceeds the configured {max_cells_allowed} used-cell limit"
                )
            headers = [
                "" if sheet.cell(1, column).value is None else str(sheet.cell(1, column).value).strip()
                for column in range(1, max_col + 1)
            ]
            rows = []
            populated_count = 0
            for row_number in range(2, max_row + 1):
                cells = [
                    "" if sheet.cell(row_number, column).value is None
                    else str(sheet.cell(row_number, column).value).strip()
                    for column in range(1, max_col + 1)
                ]
                if any(cells):
                    populated_count += 1
                    if len(rows) < intake_resolver.sample_row_limit():
                        rows.append(cells)
            hidden = sheet.sheet_state != "visible"
            tabs.append({
                "tab": sheet.title,
                "hidden": hidden,
                "headers": headers,
                "rows": rows,
                "row_count": populated_count,
            })
            sheet_states.append({"tab": sheet.title, "hidden": hidden})
            validations = getattr(sheet.data_validations, "dataValidation", [])
            for validation in validations:
                if str(validation.type or "").lower() != "list":
                    continue
                inline_options = _inline_list(validation.formula1)
                resolved_source = _resolved_list_source(
                    workbook, sheet, validation.formula1,
                )
                if inline_options is None and resolved_source is None:
                    raise ValueError(
                        f"Dropdown on tab '{sheet.title}' uses a dynamic or unsupported "
                        "list source that Parity cannot safely verify after conversion"
                    )
                dropdowns.append({
                    "tab": sheet.title,
                    "ranges": str(validation.sqref).split(),
                    "formula1": str(validation.formula1 or ""),
                    "inline_options": inline_options,
                    "resolved_options": (
                        resolved_source["options"] if resolved_source else None
                    ),
                    "source_tab": (
                        resolved_source["tab"] if resolved_source else None
                    ),
                    "source_range": (
                        resolved_source["range"] if resolved_source else None
                    ),
                    "resolved_formula": (
                        resolved_source["formula"] if resolved_source else None
                    ),
                    "allow_blank": validation.allowBlank,
                    "show_error": validation.showErrorMessage,
                    "error_style": validation.errorStyle,
                    "source_max_row": max_row,
                    "source_max_col": max_col,
                })
    finally:
        workbook.close()
        workbook_bytes.close()
    compact = {
        "tabs": [{"tab": tab["tab"], "hidden": tab["hidden"], "headers": tab["headers"],
                  "row_count": tab["row_count"]} for tab in tabs],
        "dropdowns": dropdowns,
    }
    return {
        "tabs": tabs,
        "sheet_states": sheet_states,
        "dropdowns": dropdowns,
        "fingerprint": hashlib.sha256(
            json.dumps(compact, ensure_ascii=False, sort_keys=True,
                       separators=(",", ":")).encode("utf-8")
        ).hexdigest(),
    }


def _persistable_snapshot(snapshot: dict[str, Any]) -> dict[str, Any]:
    """Keep validation contracts without retaining sampled customer cells."""
    return {
        "sheet_states": deepcopy(snapshot.get("sheet_states", [])),
        "tabs": [
            {
                "tab": tab.get("tab", ""),
                "hidden": bool(tab.get("hidden")),
                "headers": deepcopy(tab.get("headers", [])),
                "row_count": int(tab.get("row_count") or 0),
            }
            for tab in snapshot.get("tabs", [])
        ],
        "dropdowns": deepcopy(snapshot.get("dropdowns", [])),
        "fingerprint": snapshot.get("fingerprint", ""),
    }


def _quoted_tab(tab: str) -> str:
    return "'" + str(tab).replace("'", "''") + "'"


def _effective_range(token: str, max_row: int, max_col: int) -> tuple[str, int]:
    min_col, min_row, end_col, end_row = range_boundaries(token)
    min_col = min_col or 1
    end_col = end_col or 16384
    min_row = min_row or 1
    end_row = end_row or 1048576
    cell_count = (end_row - min_row + 1) * (end_col - min_col + 1)
    normalized = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(end_col)}{end_row}"
    )
    return normalized, cell_count


def verify_converted_sheet(sheet_id: str, snapshot: dict[str, Any]) -> dict[str, Any]:
    """Fail closed when Excel dropdowns or tab structure changed in conversion."""
    sheets, _ = sheets_writer._get_services()
    metadata = sheets.spreadsheets().get(
        spreadsheetId=sheet_id,
        fields="sheets.properties(sheetId,title,hidden)",
    ).execute()
    converted_states = [
        {"tab": item["properties"]["title"], "hidden": bool(item["properties"].get("hidden"))}
        for item in metadata.get("sheets", [])
    ]
    if converted_states != snapshot.get("sheet_states", []):
        raise ValueError("Google conversion changed workbook tab names, order, or visibility")
    for expected_tab in snapshot.get("tabs", []):
        values = sheets.spreadsheets().values().get(
            spreadsheetId=sheet_id,
            range=f"{_quoted_tab(expected_tab['tab'])}!1:1",
        ).execute().get("values", [])
        actual_headers = [
            str(value).strip() for value in (values[0] if values else [])
        ]
        while actual_headers and not actual_headers[-1]:
            actual_headers.pop()
        expected_headers = [str(value).strip() for value in expected_tab.get("headers", [])]
        while expected_headers and not expected_headers[-1]:
            expected_headers.pop()
        if actual_headers != expected_headers:
            raise ValueError(
                f"Google conversion changed headers on tab '{expected_tab['tab']}'"
            )

    max_ranges = max(1, int(os.getenv("WORKBOOK_MAX_DROPDOWN_RANGES", "250")))
    max_cells = max(1, int(os.getenv("WORKBOOK_MAX_DROPDOWN_CELLS", "100000")))
    dropdowns = snapshot.get("dropdowns", [])
    range_total = sum(len(item.get("ranges", [])) for item in dropdowns)
    if range_total > max_ranges:
        raise ValueError(
            f"Workbook has {range_total} dropdown ranges; safe verification limit is {max_ranges}"
        )
    verified_cells = 0
    for expected in dropdowns:
        for token in expected.get("ranges", []):
            normalized_range, cell_count = _effective_range(
                token, int(expected["source_max_row"]), int(expected["source_max_col"])
            )
            verified_cells += cell_count
            if verified_cells > max_cells:
                raise ValueError(
                    f"Workbook dropdown coverage exceeds safe verification limit of {max_cells} cells"
                )
            response = sheets.spreadsheets().get(
                spreadsheetId=sheet_id,
                ranges=[f"{_quoted_tab(expected['tab'])}!{normalized_range}"],
                includeGridData=True,
                fields="sheets(data(rowData(values(dataValidation))))",
            ).execute()
            data = (response.get("sheets") or [{}])[0].get("data") or [{}]
            row_data = data[0].get("rowData") or []
            observed = []
            for row in row_data:
                observed.extend(row.get("values") or [])
            if len(observed) < cell_count:
                raise ValueError(
                    f"Google conversion removed dropdown cells on tab '{expected['tab']}'"
                )
            inline_options = expected.get("inline_options")
            for cell in observed[:cell_count]:
                rule = cell.get("dataValidation")
                if not rule:
                    raise ValueError(
                        f"Google conversion removed a dropdown on tab '{expected['tab']}'"
                    )
                condition = rule.get("condition") or {}
                condition_type = condition.get("type")
                if inline_options is not None:
                    if condition_type != "ONE_OF_LIST":
                        raise ValueError(
                            f"Google conversion changed a dropdown rule on tab '{expected['tab']}'"
                        )
                    actual_options = [
                        str(value.get("userEnteredValue") or "")
                        for value in condition.get("values", [])
                    ]
                    if actual_options != inline_options:
                        raise ValueError(
                            f"Google conversion changed dropdown options on tab '{expected['tab']}'"
                        )
                else:
                    if condition_type not in {"ONE_OF_RANGE", "ONE_OF_LIST"}:
                        raise ValueError(
                            f"Google conversion changed a range dropdown on tab '{expected['tab']}'"
                        )
                    actual_values = [
                        str(value.get("userEnteredValue") or "")
                        for value in condition.get("values", [])
                    ]
                    resolved_options = expected.get("resolved_options")
                    if condition_type == "ONE_OF_LIST" and resolved_options is not None:
                        if actual_values != resolved_options:
                            raise ValueError(
                                f"Google conversion changed dropdown values on tab '{expected['tab']}'"
                            )
                    elif condition_type == "ONE_OF_RANGE":
                        allowed_references = {
                            _normalized_range_formula(expected.get("formula1")),
                            _normalized_range_formula(expected.get("resolved_formula")),
                        }
                        allowed_references.discard("")
                        if (
                            not actual_values
                            or _normalized_range_formula(actual_values[0])
                            not in allowed_references
                        ):
                            raise ValueError(
                                f"Google conversion changed a dropdown source range on "
                                f"tab '{expected['tab']}'"
                            )
                        source_tab = expected.get("source_tab")
                        source_range = expected.get("source_range")
                        if not source_tab or not source_range or resolved_options is None:
                            raise ValueError(
                                f"Dropdown source values on tab '{expected['tab']}' "
                                "could not be verified"
                            )
                        converted_source = sheets.spreadsheets().values().get(
                            spreadsheetId=sheet_id,
                            range=f"{_quoted_tab(source_tab)}!{source_range}",
                        ).execute().get("values", [])
                        converted_options = [
                            str(value).strip()
                            for row in converted_source
                            for value in row
                            if value is not None
                        ]
                        if converted_options != resolved_options:
                            raise ValueError(
                                f"Google conversion changed dropdown source values on "
                                f"tab '{expected['tab']}'"
                            )
                if expected.get("show_error") is not None:
                    strict_expected = bool(
                        expected.get("show_error")
                        and str(expected.get("error_style") or "stop").lower() == "stop"
                    )
                    if bool(rule.get("strict")) != strict_expected:
                        raise ValueError(
                            f"Google conversion changed dropdown reject/warn behavior on "
                            f"tab '{expected['tab']}'"
                        )
    return {"dropdown_ranges": range_total, "dropdown_cells": verified_cells}


def _trash_sheet(sheet_id: str) -> bool:
    _, drive = sheets_writer._get_services()
    try:
        drive.files().update(
            fileId=sheet_id, body={"trashed": True}, supportsAllDrives=True,
            fields="id,trashed",
        ).execute()
        return True
    except Exception:
        log.exception("Could not trash failed converted workbook %s", sheet_id)
        return False


def _convert_local_xlsx(path: str, output_name: str, run_id: str) -> str:
    from googleapiclient.http import MediaFileUpload

    folder = os.getenv("SHEET_PARENT_FOLDER_ID", "").strip()
    if not folder:
        raise RuntimeError("SHEET_PARENT_FOLDER_ID is required for workbook conversion")
    _, drive = sheets_writer._get_services()
    media = MediaFileUpload(path, mimetype=_XLSX_MIME, resumable=True)
    return drive.files().create(
        body={
            "name": Path(output_name).stem,
            "mimeType": _SHEET_MIME,
            "parents": [folder],
            "appProperties": {"parityGenerated": "true", "parityRunId": run_id},
        },
        media_body=media,
        fields="id",
        supportsAllDrives=True,
    ).execute()["id"]


def _convert_drive_xlsx(file_id: str, output_name: str, run_id: str) -> str:
    _, drive = sheets_writer._get_services()
    metadata = drive.files().get(
        fileId=file_id, fields="parents", supportsAllDrives=True
    ).execute()
    body: dict[str, Any] = {
        "name": Path(output_name).stem,
        "mimeType": _SHEET_MIME,
        "appProperties": {"parityGenerated": "true", "parityRunId": run_id},
    }
    if metadata.get("parents"):
        body["parents"] = metadata["parents"]
    return drive.files().copy(
        fileId=file_id, body=body, fields="id", supportsAllDrives=True
    ).execute()["id"]


def _download_drive_xlsx(file_id: str, destination: str) -> None:
    from googleapiclient.http import MediaIoBaseDownload

    _, drive = sheets_writer._get_services()
    request = drive.files().get_media(fileId=file_id, supportsAllDrives=True)
    with open(destination, "wb") as handle:
        downloader = MediaIoBaseDownload(handle, request)
        done = False
        while not done:
            _status, done = downloader.next_chunk()


def _selected_mappings(run: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        deepcopy(tab["mapping"])
        for tab in run.get("tabs", [])
        if tab.get("status") == "selected" and tab.get("mapping")
    ]


def _normalize_address(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _write_chunk_csv(run_id: str, index: int, items: list[dict[str, Any]]) -> str:
    fd, temp_path = tempfile.mkstemp(prefix=f"{run_id}-{index:04d}-", suffix=".csv")
    os.close(fd)
    try:
        with open(temp_path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["Address"])
            writer.writerows([[item["address"]] for item in items])
        blob = f"{RUN_ROOT}/{run_id}/chunks/{index:04d}.csv"
        upload_file(temp_path, blob)
        return blob
    finally:
        try:
            os.unlink(temp_path)
        except OSError:
            pass


def _build_queue(run: dict[str, Any]) -> dict[str, Any]:
    mappings = _selected_mappings(run)
    if not mappings:
        raise ValueError("Workbook has no selected address tabs")
    bindings = sheets_writer.read_bound_sheet_mappings(run["sheet_url"], mappings)
    from review_render import FIT_OPTIONS, HVAC_SYSTEMS, NONE_OPTION
    for binding in bindings:
        sheets_writer.ensure_review_columns(
            binding,
            binding.get("headers", []),
            HVAC_SYSTEMS + [NONE_OPTION],
            FIT_OPTIONS,
        )
    tab_counts = collections.Counter()
    grouped: collections.OrderedDict[str, dict[str, Any]] = collections.OrderedDict()
    target_order = []
    slim_bindings = []
    for binding in bindings:
        tab_counts[binding["tab"]] += len(binding.get("targets", []))
        for target in binding.get("targets", []):
            target_order.append(target["source_key"])
            normalized = _normalize_address(target["address"])
            if not normalized:
                continue
            analysis_key = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
            item = grouped.setdefault(
                analysis_key,
                {"analysis_key": analysis_key, "address": target["address"], "targets": []},
            )
            item["targets"].append({
                "source_key": target["source_key"],
                "grid_id": target["grid_id"],
                "tab": target["tab"],
                "source_row": target["source_row"],
            })
        slim_bindings.append({
            key: deepcopy(value)
            for key, value in binding.items()
            if key not in {"processing_rows", "source_rows", "targets"}
        })
    items = list(grouped.values())
    chunks = []
    size = chunk_size()
    for index, offset in enumerate(range(0, len(items), size)):
        chunk_items = items[offset:offset + size]
        chunks.append({
            "index": index,
            "status": "pending",
            "csv_blob": _write_chunk_csv(run["run_id"], index, chunk_items),
            "partial_blob": f"{RUN_ROOT}/{run['run_id']}/chunks/{index:04d}.partial.json",
            "result_blob": f"{RUN_ROOT}/{run['run_id']}/chunks/{index:04d}.result.json",
            "items": chunk_items,
        })
    for tab in run.get("tabs", []):
        tab["row_count"] = int(tab_counts.get(tab["tab"], 0))
    run["sheet_bindings"] = slim_bindings
    run["target_order"] = target_order
    run["row_count"] = len(target_order)
    run["analysis_count"] = len(items)
    run["chunks"] = chunks
    run["completed_rows"] = 0
    run["reviewed_rows"] = 0
    run["needs_attention"] = 0
    run["cost_estimate"] = _cost_estimate(len(items))
    run["status"] = (
        "approval_required"
        if len(target_order) > auto_approval_threshold()
        else "preflight"
    )
    run["confirmation_required"] = False
    run["preflight_failed"] = False
    record_metric("workbook_runs_preflighted")
    record_metric("workbook_tabs_selected", len(mappings))
    record_metric(
        "workbook_tabs_ignored",
        sum(tab.get("status") == "ignored" for tab in run.get("tabs", [])),
    )
    record_metric("workbook_rows_eligible", len(target_order))
    record_metric("workbook_unique_analyses", len(items))
    record_metric("workbook_duplicate_rows_reused", len(target_order) - len(items))
    if run["status"] == "approval_required":
        record_metric("workbook_approvals_required")
    return run


def _enqueue(run: dict[str, Any], *, approved: bool) -> dict[str, Any]:
    if run.get("usage_reserved"):
        return run
    if run.get("analysis_count", 0) <= 0:
        raise ValueError("Workbook has no usable addresses")
    payload = {
        "job_id": run["run_id"],
        "kind": "workbook",
        "workbook_manifest": _run_path(run["run_id"]),
        "total": int(run["analysis_count"]),
    }
    reserved_count = int(run.get("row_count") or run["analysis_count"])
    allowed, current_usage, message = enqueue_job_with_usage_reservation(
        run["run_id"], payload, reserved_count,
        allow_over_limit=False,
    )
    if not allowed:
        raise ValueError(message)
    run["usage_reserved"] = True
    run["reserved_address_count"] = reserved_count
    run["usage_before_reservation"] = current_usage
    run["approved_at"] = _utcnow() if approved else None
    run["status"] = "queued"
    record_metric("workbook_runs_queued")
    save_run(run)
    return run


def _convert_and_prepare(run: dict[str, Any]) -> dict[str, Any]:
    source_path = str(get_file_path(run["source_blob"]))
    sheet_id = None
    try:
        if run.get("source_drive_id"):
            sheet_id = _convert_drive_xlsx(
                run["source_drive_id"], run["source_name"], run["run_id"],
            )
        else:
            sheet_id = _convert_local_xlsx(
                source_path, run["source_name"], run["run_id"],
            )
        run["sheet_id"] = sheet_id
        run["sheet_url"] = f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit"
        run["conversion_verification"] = verify_converted_sheet(
            sheet_id, run["xlsx_snapshot"]
        )
        record_metric("workbook_conversions_verified")
        run = _build_queue(run)
        if run["status"] != "approval_required":
            run = _enqueue(run, approved=False)
        else:
            save_run(run)
        run["source_blob_deleted"] = cleanup_source_copy(run)
        save_run(run)
        return run
    except Exception as exc:
        trashed = None
        if sheet_id:
            trashed = _trash_sheet(sheet_id)
        run["status"] = "needs_attention"
        run["preflight_failed"] = True
        run["converted_copy_trashed"] = trashed
        run["error"] = str(exc)
        if sheet_id and not trashed:
            run["error"] += (
                " The failed converted copy could not be trashed automatically; "
                "an operator must remove it from Drive."
            )
        record_metric("workbook_preflight_blocks")
        save_run(run)
        return run


def prepare_local_xlsx(
    path: str | Path, filename: str, *, source_kind: str = "browser",
    source_drive_id: str | None = None,
) -> dict[str, Any]:
    if not enabled():
        raise RuntimeError("Multi-tab workbook processing is not enabled")
    if Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("Save the workbook as .xlsx and upload it again.")
    run_id = f"w-{uuid.uuid4().hex[:12]}"
    source_name = _safe_name(filename)
    source_blob = f"{RUN_ROOT}/{run_id}/source.xlsx"
    snapshot = inspect_xlsx(path)
    resolution = intake_resolver.resolve_workbook_schema(
        snapshot["tabs"], _address_variants()
    )
    upload_file(str(path), source_blob)
    run = {
        "schema_version": SCHEMA_VERSION,
        "run_id": run_id,
        "source_kind": source_kind,
        "source_name": source_name,
        "source_blob": source_blob,
        "source_drive_id": source_drive_id,
        "source_fingerprint": snapshot["fingerprint"],
        "xlsx_snapshot": _persistable_snapshot(snapshot),
        "tabs": resolution["tabs"],
        "mapping_fingerprint": resolution["fingerprint"],
        "status": "preflight",
        "confirmation_required": resolution["status"] != "ready",
        "created_at": _utcnow(),
        "updated_at": _utcnow(),
    }
    save_run(run)
    if run["confirmation_required"]:
        return run
    return _convert_and_prepare(run)


def _read_csv_rows(path: str | Path) -> list[list[str]]:
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(path, "r", encoding=encoding, newline="") as handle:
                sample = handle.read(4096)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                rows = []
                max_rows = max(
                    1, int(os.getenv("WORKBOOK_MAX_ROWS_PER_TAB", "100000"))
                )
                max_cols = max(1, int(os.getenv("WORKBOOK_MAX_COLUMNS", "500")))
                max_cells = max(
                    1, int(os.getenv("WORKBOOK_MAX_USED_CELLS", "5000000"))
                )
                used_cells = 0
                for row in csv.reader(handle, dialect):
                    if len(rows) >= max_rows:
                        raise ValueError(
                            f"CSV exceeds the configured {max_rows} row limit"
                        )
                    if len(row) > max_cols:
                        raise ValueError(
                            f"CSV exceeds the configured {max_cols} column limit"
                        )
                    used_cells += len(row)
                    if used_cells > max_cells:
                        raise ValueError(
                            f"CSV exceeds the configured {max_cells} used-cell limit"
                        )
                    rows.append(list(row))
                return rows
        except (OSError, UnicodeError, csv.Error) as exc:
            last_error = exc
    raise ValueError(f"Could not read CSV: {last_error}")


def prepare_local_csv(
    path: str | Path, filename: str, *, source_kind: str = "browser",
) -> dict[str, Any]:
    """Route a CSV through the same engine as a one-tab, no-dropdown workbook."""
    rows = _read_csv_rows(path)
    if not rows or not any(any(str(cell).strip() for cell in row) for row in rows):
        raise ValueError("CSV has no usable rows")
    fd, temp_xlsx = tempfile.mkstemp(prefix="parity-csv-", suffix=".xlsx")
    os.close(fd)
    try:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Uploaded file")
        for row in rows:
            sheet.append(row)
        workbook.save(temp_xlsx)
        xlsx_name = f"{Path(filename).stem or 'upload'}.xlsx"
        run = prepare_local_xlsx(
            temp_xlsx, xlsx_name, source_kind=source_kind,
        )
        run["source_name"] = _safe_name(filename)
        save_run(run)
        return run
    finally:
        try:
            os.unlink(temp_xlsx)
        except OSError:
            pass


def prepare_local_file(
    path: str | Path, filename: str, *, source_kind: str = "browser",
) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xls":
        raise ValueError("Save the legacy Excel workbook as .xlsx and upload it again.")
    if suffix == ".xlsx":
        return prepare_local_xlsx(path, filename, source_kind=source_kind)
    if suffix == ".csv":
        return prepare_local_csv(path, filename, source_kind=source_kind)
    raise ValueError("Upload must be an .xlsx or .csv file.")


def prepare_drive_xlsx(file_id: str, filename: str) -> dict[str, Any]:
    fd, temp_path = tempfile.mkstemp(prefix="parity-drive-", suffix=".xlsx")
    os.close(fd)
    try:
        _download_drive_xlsx(file_id, temp_path)
        return prepare_local_xlsx(
            temp_path, filename, source_kind="drive_inbox", source_drive_id=file_id
        )
    finally:
        try:
            os.unlink(temp_path)
        except OSError:
            pass


def _google_sheet_tabs(sheet_url: str) -> tuple[list[dict[str, Any]], str]:
    inspected = sheets_writer.inspect_bound_sheet(sheet_url)
    return inspected["tabs"], inspected["title"]


def prepare_google_sheet(
    sheet_url: str, *, source_name: str = "Google Sheet",
    source_kind: str = "browser_sheet",
) -> dict[str, Any]:
    if not enabled():
        raise RuntimeError("Multi-tab workbook processing is not enabled")
    tabs, title = _google_sheet_tabs(sheet_url)
    resolution = intake_resolver.resolve_workbook_schema(tabs, _address_variants())
    run_id = f"w-{uuid.uuid4().hex[:12]}"
    run = {
        "schema_version": SCHEMA_VERSION,
        "run_id": run_id,
        "source_kind": source_kind,
        "source_name": source_name or title or "Google Sheet",
        "sheet_id": sheets_writer.sheet_id_from_url(sheet_url),
        "sheet_url": sheet_url,
        "tabs": resolution["tabs"],
        "mapping_fingerprint": resolution["fingerprint"],
        "status": "preflight",
        "confirmation_required": resolution["status"] != "ready",
        "created_at": _utcnow(),
        "updated_at": _utcnow(),
    }
    save_run(run)
    if run["confirmation_required"]:
        return run
    try:
        run = _build_queue(run)
        if run["status"] != "approval_required":
            run = _enqueue(run, approved=False)
        else:
            save_run(run)
    except Exception as exc:
        run["status"] = "needs_attention"
        run["preflight_failed"] = True
        run["error"] = str(exc)
        save_run(run)
    return run


def _address_variants() -> list[str]:
    from tasks_local import ADDRESS_VARIANTS
    return list(ADDRESS_VARIANTS)


def _source_tabs_for_confirmation(run: dict[str, Any]) -> list[dict[str, Any]]:
    if run.get("source_blob"):
        return inspect_xlsx(get_file_path(run["source_blob"]))["tabs"]
    tabs, _title = _google_sheet_tabs(run["sheet_url"])
    return tabs


def confirm_mappings(run_id: str, selections: list[dict[str, Any]]) -> dict[str, Any]:
    run = load_run(run_id)
    if not run or run.get("status") != "preflight" or not run.get("confirmation_required"):
        raise ValueError("Workbook setup is not available")
    source_tabs = _source_tabs_for_confirmation(run)
    by_name = {str(item.get("tab") or ""): item for item in selections}
    updated = []
    for item in run.get("tabs", []):
        if item.get("status") in {"selected", "ignored"}:
            updated.append(item)
            continue
        choice = by_name.get(item.get("tab", ""))
        if not choice:
            raise ValueError(f"Choose how to handle tab '{item.get('tab', '')}'")
        classification = str(choice.get("classification") or "").lower()
        if classification == "ignore":
            updated.append({**item, "status": "ignored",
                            "reason": "ignored by confirmed workbook setup",
                            "mapping": None})
            continue
        if classification != "address":
            raise ValueError(f"Invalid classification for tab '{item.get('tab', '')}'")
        candidate = {
            "tab": item["tab"],
            "address_column": choice.get("address_column"),
            "city_column": choice.get("city_column"),
            "state_column": choice.get("state_column"),
            "zip_column": choice.get("zip_column"),
            "confidence": 1.0,
        }
        valid, reason, mapping = intake_resolver.validate_mapping(source_tabs, candidate)
        if not valid or mapping is None:
            raise ValueError(f"Tab '{item['tab']}' mapping is invalid: {reason}")
        updated.append({**item, "status": "selected", "mapping": mapping,
                        "reason": "confirmed address mapping"})
    run["tabs"] = updated
    run["status"] = "preflight"
    run["confirmation_required"] = False
    run.pop("error", None)
    save_run(run)
    if run.get("source_blob"):
        return _convert_and_prepare(run)
    try:
        run = _build_queue(run)
        if run["status"] != "approval_required":
            return _enqueue(run, approved=False)
        return save_run(run)
    except Exception as exc:
        run["status"] = "needs_attention"
        run["preflight_failed"] = True
        run["error"] = str(exc)
        return save_run(run)


def approve(run_id: str) -> dict[str, Any]:
    run = load_run(run_id)
    if not run or run.get("status") != "approval_required":
        raise ValueError("Workbook is not waiting for approval")
    return _enqueue(run, approved=True)


def retry(run_id: str) -> dict[str, Any]:
    """Resume a failed analysis job from its durable row checkpoints."""
    run = load_run(run_id)
    if not run:
        raise ValueError("Workbook run was not found")
    if run.get("preflight_failed"):
        raise ValueError("Preflight failures must be corrected and uploaded again")
    if run.get("review_url"):
        raise ValueError("Use the row rerun action for a workbook already in review")
    if run.get("status") != "needs_attention" or not run.get("usage_reserved"):
        raise ValueError("Workbook run is not eligible for checkpoint retry")
    job = get_job_status(run_id)
    if not job or job.get("status") != "failed":
        raise ValueError("Workbook job is not in a retryable failed state")

    def mark_queued(current):
        current["status"] = "queued"
        current["retry_requested_at"] = _utcnow()
        current.pop("error", None)

    mutate_run(run_id, mark_queued)
    queued, message = requeue_failed_workbook_job(run_id)
    if not queued:
        def restore_attention(current):
            current["status"] = "needs_attention"
            current["error"] = message
        mutate_run(run_id, restore_attention)
        raise ValueError(message)
    record_metric("workbook_checkpoint_retries")
    return load_run(run_id) or run


def update_review_progress(
    run_id: str, reviewed: int, total: int, writeback_failures: int = 0,
    needs_attention: int | None = None,
) -> None:
    def updater(run):
        run["reviewed_rows"] = int(reviewed)
        run["writeback_failures"] = int(writeback_failures)
        if needs_attention is not None:
            run["needs_attention"] = int(needs_attention)
        run["status"] = (
            "complete"
            if total and reviewed >= total and not writeback_failures
            else (
                "needs_attention"
                if run.get("needs_attention") or writeback_failures
                else "review_open"
            )
        )
        if run["status"] == "complete":
            run["completed_at"] = _utcnow()
    mutate_run(run_id, updater)


def update_analysis_result(
    run_id: str, *, review_url: str, completed_rows: int,
    needs_attention: int,
) -> None:
    def updater(run):
        run["review_url"] = review_url
        run["completed_rows"] = int(completed_rows)
        run["needs_attention"] = int(needs_attention)
        run["status"] = "needs_attention" if needs_attention else "review_open"
    mutate_run(run_id, updater)


def chunk_partial_path(run_id: str, index: int) -> str:
    run = load_run(run_id) or {}
    for chunk in run.get("chunks", []):
        if int(chunk.get("index", -1)) == int(index):
            return chunk["partial_blob"]
    raise KeyError(index)


def mark_chunk_complete(run_id: str, index: int) -> None:
    def updater(run):
        for chunk in run.get("chunks", []):
            if int(chunk.get("index", -1)) == int(index):
                chunk["status"] = "complete"
                break
    mutate_run(run_id, updater)


def cleanup_source_copy(run: dict[str, Any]) -> bool:
    """Remove only the app's private upload copy after a successful conversion."""
    if run.get("source_kind") in {"browser", "api", "drive_inbox"} and run.get("source_blob"):
        try:
            delete_file(run["source_blob"])
            return True
        except Exception:
            log.warning("Could not remove private source copy for %s", run.get("run_id"))
            return False
    return False
